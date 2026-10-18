from pathlib import Path
import json,zipfile,xml.etree.ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parents[1]
wb=Workbook();wb.remove(wb.active);cached={}
money='"$"#,##0';precise='"$"#,##0.00'
navy='173D4A';teal='087F8C'
def sheet(name,title,note,headers,widths=None):
 ws=wb.create_sheet(name);ws.append([title]);ws.append([note]);ws.append([]);ws.append(headers)
 ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers));ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
 ws.row_dimensions[1].height=30;ws.row_dimensions[2].height=55
 ws['A1'].font=Font(name='Calibri',size=17,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor=navy)
 ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws['A2'].font=Font(size=10,color='475569')
 for c in ws[4]:c.fill=PatternFill('solid',fgColor=teal);c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[4].height=32;ws.freeze_panes='B5';ws.sheet_view.showGridLines=False
 for i in range(1,len(headers)+1):ws.column_dimensions[chr(64+i)].width=(widths[i-1] if widths else 24)
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:4';return ws

def formula(ws,cell,expr,value,fmt=money):
 ws[cell]=expr;ws[cell].number_format=fmt;cached[(ws.title,cell)]=value

ws=sheet('Overview','Budget levels — 17 October 2026','USD planning model; no hardware owned. Equipment estimates, professional time and human validation are separate. No purchase, flash, clinical approval or human-use authorization.',['Decision','Recommended interpretation'],[34,115])
for row in [
 ('First bench','A3: $12k–$25k equipment/setup; illustrative cart $17,203.20 including hardware reserve.'),
 ('Professionally run first phase','Illustrative $77,920 before tax/freight, using one program reserve rather than double-counting the separate hardware reserve.'),
 ('Human-connected engineering','Choose ONE cumulative B tier. B includes an approximately $20k baseline bench allowance; do not add the entire bench budget again.'),
 ('Physiological validation','Choose ONE overall C program tier. Participant counts are planning scenarios, not powered designs. Custom device engineering is separate.'),
 ('Cash vs time','Bench initial cash excludes labor. Unpaid owner hours have economic cost. The Labor and TCO sheet exposes professional-time assumptions.'),
 ('Recurring costs','Bench annual allowances cover cash upkeep/access after year one, not recurring professional staff. Five-year cash = setup + four upkeep years.'),
 ('Price confidence','Only named price anchors are observed listings. All wider budgets and rates are original planning assumptions, not supplier quotes.'),
 ('Source of truth','Read the companion report for scope, exclusions, sources and safety gates. Current software remains nominal 50 kHz only.'),
 ('Safety','Track A is passive-only. Neither an evaluation board nor the link-check ELF is cleared for human connection or flashing.'),
 ('Edit inputs','Blue entries are editable assumptions. Excel/compatible software recalculates formulas on opening; initial formula caches are provided.')]:ws.append(row)

ass=sheet('Assumptions','Editable planning assumptions','These are original scenario inputs, NOT wage surveys, facility quotes, powered study designs or guaranteed project costs.',['Input','Value','Unit / purpose'],[39,22,100])
inputs=[('Professional rate — low',150,'USD/hour'),('Professional rate — high',250,'USD/hour'),('Phase-1 rate',200,'USD/hour'),('Hardware reserve',.20,'fraction; cart only'),('Program reserve',.25,'fraction; phase-1 or study scenario'),('Participants in study example',300,'illustrative; not powered'),('Visits per participant',2,'illustrative'),('Cost per visit',800,'USD; fully burdened assumption, define inclusions'),('Study overhead',200000,'USD beyond fully burdened visits; avoid duplicated costs'),('Loaded FTE cost',240000,'USD/person-year assumption'),('FTE example headcount',4,'people'),('FTE example years',1,'years'),('TCO years',5,'initial setup includes initial-year services'),('Phase-1 professional hours',200,'one-time person-hours'),('Independent DUT verification',8000,'USD; distinct from initial standard/instrument calibration')]
for label,value,note in inputs:ass.append([label,value,note]);ass.cell(ass.max_row,2).font=Font(color='2563EB')
# Mapping of assumption rows is explicit and used in formulas.
A={label:row for row,(label,_,_) in enumerate(inputs,start=5)}
def ref(name):return f"Assumptions!$B${A[name]}"
for name in ['Hardware reserve','Program reserve']:ass.cell(A[name],2).number_format='0%'

bench=sheet('Bench','Track A — alternative passive-only setups','Initial cash includes equipment/fixtures and specified initial services, excludes labor/tax/freight. Annual upkeep is modeled for years AFTER the initial year. Do not sum tiers.',['Tier','Setup low','Setup high','Annual low','Annual high','Scope'],[31,18,18,18,18,90])
benchrows=[('A1 Learning',1000,3000,200,800,'Bring-up, logging and simple passive checks; not traceable phase accuracy.'),('A2 Shared reference',3000,8000,1000,4000,'Defined lab/rental/service campaign; rates and proper fixture require quotes.'),('A3 Owned reference',12000,25000,1500,4500,'Appropriate LCR, correct fixture, calibration/standards, scope, DMM and DUT.'),('A4 Premium characterization',35000,100000,3000,10000,'Advanced analyzer/workflow, programmable passive fixtures and temperature work.'),('A5 Automated lab',100000,300000,8000,30000,'Automated regression/environment/unit-to-unit verification; upper bounds can be exceeded.')]
for row in benchrows:bench.append(row)
for row in bench.iter_rows(min_row=5,min_col=2,max_col=5):
 for c in row:c.number_format=money;c.font=Font(color='2563EB')

lab=sheet('Labor and TCO','Bench labor and five-year cash scenarios','Professional work is one-time setup/characterization effort. Five-year cash excludes labor; professional staff for ongoing operations are NOT included. Tiers are alternatives.',['Tier','Hours low','Hours high','Labor low','Labor high','Initial + labor low','Initial + labor high','Five-year cash low','Five-year cash high'],[31,15,15,21,21,23,23,23,23])
for row,hours in zip(range(5,10),[(40,120),(80,200),(150,400),(400,1000),(1500,4000)]):
 base=benchrows[row-5];lo,hi=hours;lab.append([base[0],lo,hi]);lab.cell(row,2).font=Font(color='2563EB');lab.cell(row,3).font=Font(color='2563EB')
 formula(lab,f'D{row}',f'=B{row}*{ref("Professional rate — low")}',lo*150)
 formula(lab,f'E{row}',f'=C{row}*{ref("Professional rate — high")}',hi*250)
 formula(lab,f'F{row}',f'=Bench!B{row}+D{row}',base[1]+lo*150)
 formula(lab,f'G{row}',f'=Bench!C{row}+E{row}',base[2]+hi*250)
 formula(lab,f'H{row}',f'=Bench!B{row}+({ref("TCO years")}-1)*Bench!D{row}',base[1]+4*base[3])
 formula(lab,f'I{row}',f'=Bench!C{row}+({ref("TCO years")}-1)*Bench!E{row}',base[2]+4*base[4])

eng=sheet('Engineering','Track B — cumulative custom-device engineering','Includes an approximately $20k baseline bench allowance and engineering/compliance. Excludes large human studies, mass-production inventory and ongoing commercial operations. Choose ONE tier; do not sum.',['Tier','Budget low','Budget high','Window','Milestone / limits'],[40,22,22,25,100])
for row in [('B1 Passive custom instrument',100000,300000,'3–6 months','Custom electronics, raw export, PCB iterations and passive verification; no human connection.'),('B2 Reviewed research instrument',300000,1000000,'6–12 months','Complete safety architecture, specialist testing/documentation; human work requires appropriate approvals.'),('B3 Segmental multifrequency platform',750000,2500000,'12–24 months','Characterized switching, contact QC, timing, mechanical/electrode/scale integration; not a validation result.'),('B4 Productized composition platform',2000000,8000000,'18–36+ months','Manufacturing, quality, software lifecycle, usability/cybersecurity, regulatory program for chosen claims; can exceed range.')]:eng.append(row)

val=sheet('Validation','Track C — physiological evidence programs','Budgets through the stated milestone, including appropriate pilot/prep. Excludes custom device engineering and imaging-facility purchase. Counts are illustrative, not powered sample-size recommendations.',['Tier','Budget low','Budget high','Participants low','Participants high','Scope'],[40,22,22,21,21,100])
for row in [('C1 Repeatability/feasibility',50000,200000,20,50,'Repeated visits/comparators after approvals. Workflow and preliminary evidence, not universal accuracy.'),('C2 Focused composition study',500000,2000000,300,600,'DXA/reference work, participant-level model holdout, specific population and claims.'),('C3 Broader/multi-site validation',2000000,8000000,1000,3000,'Independent/subgroup/longitudinal evidence and selected MRI/4-compartment/fluid references; may exceed range.')]:val.append(row)

combo=sheet('Combined scenarios','Combined program envelopes — not quotes','Engineering B includes baseline bench allowance. Validation C excludes engineering. These rows are alternative scenarios; do NOT sum them, or add every earlier tier. Broader multi-modality scope can exceed all of these.',['Scenario','Engineering choice','Validation choice','Combined low','Combined high','Interpretation'],[36,27,27,24,24,95])
for row,er,vr in [(5,6,5),(6,7,6),(7,8,7)]:
 combo.append([['Research feasibility','Focused research platform','Product + broader validation'][row-5],eng.cell(er,1).value,val.cell(vr,1).value])
 formula(combo,f'D{row}',f'=Engineering!B{er}+Validation!B{vr}',eng.cell(er,2).value+val.cell(vr,2).value)
 formula(combo,f'E{row}',f'=Engineering!C{er}+Validation!C{vr}',eng.cell(er,3).value+val.cell(vr,3).value)
 combo.cell(row,6,'No claim of certification, success, clearance or applicability to a person; major scope changes need re-estimation.')

cart=sheet('A3 cart','Illustrative A3 equipment/setup cart','Only the $5,186 LCR line is a retrieved current listed-price anchor. Other entries are allowances; quotes, exact packing list, proper probe/fixture and uncertainty scope are required. Tax, freight, new computer and labor excluded.',['Item','Quantity','Unit allowance','Extended','Basis / qualification'],[44,15,22,22,100])
cartrows=[('Kit / optional spare allowance',1,1100,'Planning allowance, not a requirement to buy two kits. Confirm controller/Z-Test/cable contents.'),('Hioki IM3536',1,5186,'Detailed TEquipment listing; probe/fixture separate; stock expected Oct 21, not verified on hand.'),('Correct probe/interface/fixture',1,1000,'Planning allowance; match the reference plane and passive DUT, not an arbitrary SMD accessory.'),('Standards and initial AC calibration',1,2000,'Planning allowance; frequency/phase/uncertainty scope must be applicable.'),('DMM',1,750,'Planning allowance; not an AC phase reference.'),('Oscilloscope',1,1800,'Planning allowance; debug and timing, not a substitute for LCR reference.'),('Bench supply',1,500,'Planning allowance; passive bench only.'),('Additional passive fixture work',1,1500,'Planning allowance distinct from the basic probe interface.'),('ESD/cables/local logging',1,500,'Planning allowance; manufacturer-approved connection route only.')]
for r,(name,qty,unit,note) in enumerate(cartrows,start=5):
 cart.append([name,qty,unit,None,note]);formula(cart,f'D{r}',f'=B{r}*C{r}',qty*unit,precise);cart.cell(r,3).number_format=precise;cart.cell(r,3).font=Font(color='2563EB')
cart.append(['Subtotal']);formula(cart,'D14','=SUM(D5:D13)',14336,precise)
cart.append(['Hardware reserve']);formula(cart,'D15',f'=D14*{ref("Hardware reserve")}',2867.2,precise)
cart.append(['Equipment allocation']);formula(cart,'D16','=D14+D15',17203.2,precise)
cart['E6'].hyperlink='https://www.tequipment.net/Hioki/IM3536/LCR-Meters/'

phase=sheet('Phase 1','Professionally run first-phase example','Uses cart SUBTOTAL without its separate reserve, then applies ONE program reserve. Engineering and independent verification are modeled allowances; no quotation or purchase authorization.',['Component','Amount','Meaning'],[53,27,105])
phase.append(['Equipment/setup before reserve']);formula(phase,'B5',"='A3 cart'!D14",14336,precise)
phase.append(['Professional work']);formula(phase,'B6',f'={ref("Phase-1 professional hours")}*{ref("Phase-1 rate")}',40000,precise)
phase.append(['Independent DUT/fixture verification']);formula(phase,'B7',f'={ref("Independent DUT verification")}',8000,precise)
phase.append(['Subtotal']);formula(phase,'B8','=SUM(B5:B7)',62336,precise)
phase.append(['Program reserve']);formula(phase,'B9',f'=B8*{ref("Program reserve")}',15584,precise)
phase.append(['Total']);formula(phase,'B10','=B8+B9',77920,precise)
phase['C5']='Does not also include the cart’s 20% reserve.';phase['C7']='Different scope from initial reference/standards calibration; obtain an actual service proposal.'
phase.append([]);phase.append(['Study example — not a protocol'])
phase.append(['Fully burdened visit costs']);formula(phase,'B13',f'={ref("Participants in study example")}*{ref("Visits per participant")}*{ref("Cost per visit")}',480000,precise)
phase.append(['Additional study overhead']);formula(phase,'B14',f'={ref("Study overhead")}',200000,precise)
phase.append(['Study example including reserve']);formula(phase,'B15',f'=(B13+B14)*(1+{ref("Program reserve")})',850000,precise)
phase.append([]);phase.append(['FTE example — not added to consultant hours'])
phase.append(['Modeled team cost']);formula(phase,'B18',f'={ref("Loaded FTE cost")}*{ref("FTE example headcount")}*{ref("FTE example years")}',960000,precise)

own=sheet('Ownership','Ownership, subscriptions and evidence gates','Price observations are not contracts. Recurring calibration/consumables can remain real costs in an entirely offline system. No legacy Withings account entitlements assumed.',['Check','Required answer / evidence'],[42,125])
for row in [('Offline operation','Acquisition/local viewing after cancellation; installer/firmware access; required accounts and network documented.'),('Raw data','Sample per-path/per-frequency R/X/IQ export plus quality flags/config/calibration. CSV or a USB port alone is insufficient.'),('License scope','Perpetual vs annual/per-seat/per-test; API/export rights; used-unit transfer; cancellation behavior.'),('Recurring costs','Separate compulsory acquisition access from optional support, calibration, electrodes and financial leases.'),('Hioki price anchor','IM3536 $5,186 detailed TEquipment page; separate fixture/probe; expected Oct 21 stock, quote needed.'),('AD5940 kit anchor','Mouser detailed page: $497.15; possible 8% US tariff; distributor end-of-stock warning. Packing list must be confirmed.'),('Keysight current family','E4980BL US page: $7,391 / $10,839 / $14,264 for 300k/500k/1MHz configurations; options/fixture/quote required. E4980AL discontinued/currently supported.'),('MFIA anchor','Manufacturer displays $14,490 region-specific/ex-works. 500kHz base vs 5MHz upgrade/configuration must be quoted; not an all-options 5MHz price.'),('Rental caution','$649/month IM3536 appears in search only; detailed rental fetch did not substantiate rate. Do not book or budget it as confirmed.'),('Safety boundary','Evaluation board passive-only. Custom patient-connected work needs complete-system safety review and required approvals. Battery/USB isolation alone is not qualification.'),('Validation boundary','Models, appropriate comparators, independent participants/sites and honest errors are required; maximum frequency or displayed metric count does not establish validity.')]:own.append(row)

for ws in wb:
 ws.auto_filter.ref=f'A4:{ws.cell(ws.max_row,ws.max_column).coordinate}'
 for row in ws.iter_rows(min_row=5):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if cell.row%2:cell.fill=PatternFill('solid',fgColor='F1F6F8')
  ws.row_dimensions[row[0].row].height=52 if ws.title not in ['Assumptions','A3 cart','Phase 1'] else 39
 for row in ws.iter_rows(min_row=5,min_col=2,max_col=min(3,ws.max_column)):
  for cell in row:
   if ws.title in ['Engineering','Validation']:cell.number_format=money;cell.font=Font(color='2563EB')
 ws.sheet_properties.outlinePr.summaryRight=True
 ws.oddFooter.center.text='Planning estimates • 17 October 2026 • No human-use authorization'
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
path=ROOT/'Body-Composition-Budget-Levels.xlsx';wb.save(path)
# Cache initial numeric formula results for preview; Excel recalculates after edits.
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ET.register_namespace('',ns)
with zipfile.ZipFile(path) as z:payload={name:z.read(name) for name in z.namelist()}
for i,ws in enumerate(wb.worksheets,start=1):
 name=f'xl/worksheets/sheet{i}.xml';tree=ET.fromstring(payload[name])
 for c in tree.findall(f'.//{{{ns}}}c'):
  key=(ws.title,c.get('r'))
  if key in cached:
   v=c.find(f'{{{ns}}}v')
   if v is None:v=ET.SubElement(c,f'{{{ns}}}v')
   v.text=str(cached[key])
 payload[name]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in payload.items():z.writestr(name,data)
values=load_workbook(path,data_only=True)
assert values['A3 cart']['D16'].value==17203.2
assert values['Phase 1']['B10'].value==77920
assert values['Phase 1']['B15'].value==850000
assert values['Phase 1']['B18'].value==960000
assert values['Combined scenarios']['D5'].value==350000
assert values['Combined scenarios']['E7'].value==16000000
assert all(values[name][cell].value==v for (name,cell),v in cached.items())
(ROOT/'budget-levels/workbook-check.json').write_text(json.dumps({'formula_cells_checked':len(cached),'sheets':len(wb.worksheets),'initial_cached_results_verified':True,'recalculate_on_open':True,'equipment_example_usd':17203.2,'phase1_example_usd':77920,'study_example_usd':850000,'fte_example_usd':960000,'note':'Budgets/rates are planning assumptions, not quotations; cached values reflect initial assumptions.'},indent=2)+'\n')
print('Workbook saved;',len(wb.worksheets),'sheets;',len(cached),'formula caches checked.')
